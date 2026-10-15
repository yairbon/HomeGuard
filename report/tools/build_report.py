"""Builds the final report (.docx) from the three markdown parts.

    python build_report.py <report_dir> <figs_dir> <out_docx>

* Layout follows the department rules: Arial 12, line spacing 1.5, RTL,
  US-Letter with 1" margins (same as the department's project template),
  page numbers from page 2, figure captions under figures, table captions
  above tables, hierarchical numbered headings (the numbers are in the text).
* Every "[הכנס ...]" placeholder whose picture exists in <figs_dir>
  (figN.png / codeN.png) is replaced by that picture; the rest stay as
  yellow-highlighted boxes so they are easy to find (Ctrl+F "[הכנס").
* Pipeline: markdown -> pandoc (math becomes native Word equations) ->
  python-docx post-processing (RTL, fonts, tables, images, cover, footer).
"""
import copy
import math
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor
from PIL import Image

REPORT_DIR = Path(sys.argv[1])
FIGS = Path(sys.argv[2])
OUT = Path(sys.argv[3])
FONT = "Arial"
HEBREW = re.compile(r"[֐-׿]")
PLACEHOLDER = re.compile(r"\[(הכנס|להשלים|לאמת|הוסף קישור|תאריך)|dd/mm/yy")
MAX_W_IN, MAX_H_IN = 6.3, 8.0

PARTS = ["01_cover_abstract_intro.md", "02_design_implementation.md", "03_tests_summary_bom_refs.md"]

# --------------------------------------------------------------------------
# 1. Markdown preprocessing
# --------------------------------------------------------------------------


def figure_files(kind: str, num: str):
    """Returns the list of picture paths for 'איור N' / 'קוד N', or []."""
    stem = ("fig" if kind == "איור" else "code") + num
    single = FIGS / f"{stem}.png"
    if single.exists():
        return [single]
    parts = sorted(FIGS.glob(f"{stem}_*.png"))
    return parts


def preprocess(text: str) -> str:
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    text = re.sub(r'<div align="center">.*?</div>', "", text, count=1, flags=re.S)  # cover (built separately)
    text = re.sub(r'<div style="page-break-after: always;"></div>', "", text)
    lines = text.split("\n")
    lines = [ln for ln in lines if ln.strip() != "---"]

    # group consecutive blockquote lines
    out, i = [], 0
    while i < len(lines):
        if lines[i].startswith(">"):
            j = i
            group = []
            while j < len(lines) and lines[j].startswith(">"):
                group.append(re.sub(r"^>\s?", "", lines[j]))
                j += 1
            paras, cur = [], []
            for g in group:
                if g.strip() == "":
                    if cur:
                        paras.append(" ".join(cur))
                        cur = []
                else:
                    cur.append(g.strip())
            if cur:
                paras.append(" ".join(cur))
            out.append("")
            handled = False
            if len(paras) == 2 and "[הכנס" in paras[0]:
                m = re.match(r"\*\*(איור|קוד) (\d+): (.*)\*\*$", paras[1])
                if m:
                    files = figure_files(m.group(1), m.group(2))
                    if files:
                        cap = f"{m.group(1)} {m.group(2)}: {m.group(3)}"
                        out.append("@@IMG|" + "|".join(str(f) for f in files) + "|" + cap + "@@")
                        handled = True
            if not handled:
                for p in paras:
                    out.append(p)
                    out.append("")
            out.append("")
            i = j
        else:
            out.append(lines[i])
            i += 1
    return "\n".join(out)


def build_master(tmp: Path) -> Path:
    md = "\n\n".join(preprocess((REPORT_DIR / p).read_text(encoding="utf-8")) for p in PARTS)
    md = md.replace("‏", "")
    path = tmp / "master.md"
    path.write_text(md, encoding="utf-8")
    return path


# --------------------------------------------------------------------------
# 2. Reference document (styles)
# --------------------------------------------------------------------------


def insert_ordered(parent, child, successors):
    for tag in successors:
        found = parent.find(qn(tag))
        if found is not None:
            found.addprevious(child)
            return
    parent.append(child)


PPR_AFTER_BIDI = ["w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                  "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                  "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"]


def set_bidi(ppr):
    if ppr.find(qn("w:bidi")) is None:
        insert_ordered(ppr, OxmlElement("w:bidi"), PPR_AFTER_BIDI)


def set_fonts(rpr, size_half_pts=None, bold=None):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


def make_reference(path: Path):
    data = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], capture_output=True, check=True).stdout
    path.write_bytes(data)
    doc = Document(str(path))
    st = doc.styles
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Inches(1))

    def style_pf(style, size=12, bold=None, italic=None, spacing=1.5, before=0, after=6, align=None,
                 keep_next=None, color=RGBColor(0, 0, 0), bidi=True):
        el = style.element
        ppr = el.get_or_add_pPr()
        if bidi:
            set_bidi(ppr)
        rpr = el.get_or_add_rPr()
        set_fonts(rpr)
        style.font.size = Pt(size)
        # complex-script size/bold/italic so Hebrew follows the same formatting
        for tag, val in (("w:szCs", str(int(size * 2))),):
            e = rpr.find(qn(tag))
            if e is None:
                e = OxmlElement(tag)
                rpr.append(e)
            e.set(qn("w:val"), val)
        if bold is not None:
            style.font.bold = bold
            e = rpr.find(qn("w:bCs"))
            if bold and e is None:
                rpr.append(OxmlElement("w:bCs"))
        if italic is not None:
            style.font.italic = italic
        if color is not None:
            style.font.color.rgb = color
        pf = style.paragraph_format
        pf.line_spacing = spacing
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        if align is not None:
            pf.alignment = align
        if keep_next is not None:
            pf.keep_with_next = keep_next

    names = {s.name: s for s in st}
    for n in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text", "Table Caption", "Definition",
              "Definition Term", "Footnote Text", "Bibliography"):
        if n in names:
            style_pf(names[n], 12, spacing=1.5, after=6,
                     align=WD_ALIGN_PARAGRAPH.JUSTIFY if n in ("Body Text", "First Paragraph") else None)
    # table cell text: single spacing, 11pt
    if "Compact" in names:
        style_pf(names["Compact"], 11, spacing=1.0, after=2, align=None)
        names["Compact"].base_style = names["Normal"]
    sizes = {"Heading 1": (18, 18, 8), "Heading 2": (14, 14, 6), "Heading 3": (13, 10, 4),
             "Heading 4": (12, 8, 4), "Heading 5": (12, 6, 3), "Heading 6": (12, 6, 3)}
    for n, (sz, before, after) in sizes.items():
        if n in names:
            style_pf(names[n], sz, bold=True, italic=False, spacing=1.15, before=before, after=after,
                     align=None, keep_next=True)
    for n in ("Caption", "Image Caption"):
        if n in names:
            style_pf(names[n], 11, bold=False, italic=False, spacing=1.0, before=2, after=10,
                     align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x33, 0x33, 0x33))
    if "Verbatim Char" in names:
        vc = names["Verbatim Char"]
        vc.font.name = "Courier New"
        vc.font.size = Pt(10.5)
        rf = vc.element.get_or_add_rPr().find(qn("w:rFonts"))
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), "Courier New")
    if "Hyperlink" in names:
        names["Hyperlink"].font.color.rgb = RGBColor(0x1A, 0x56, 0xC0)
    doc.save(str(path))


# --------------------------------------------------------------------------
# 3. Post-processing
# --------------------------------------------------------------------------


def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def cell_borders(tcpr, color="999999", sz="4"):
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), sz)
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tcpr.append(b)


def add_run_props(run_el):
    """bCs/iCs/szCs next to b/i/sz so Hebrew text is bold/italic/sized like Latin text."""
    rpr = run_el.find(qn("w:rPr"))
    if rpr is None:
        return
    pairs = (("w:b", "w:bCs"), ("w:i", "w:iCs"))
    for a, b in pairs:
        if rpr.find(qn(a)) is not None and rpr.find(qn(b)) is None:
            rpr.find(qn(a)).addnext(OxmlElement(b))
    sz = rpr.find(qn("w:sz"))
    if sz is not None and rpr.find(qn("w:szCs")) is None:
        e = OxmlElement("w:szCs")
        e.set(qn("w:val"), sz.get(qn("w:val")))
        sz.addnext(e)


def highlight_run(run):
    rpr = run._r.get_or_add_rPr()
    if rpr.find(qn("w:highlight")) is None:
        h = OxmlElement("w:highlight")
        h.set(qn("w:val"), "yellow")
        # schema order: ... color, spacing, w, kern, position, sz, szCs, highlight, u ...
        anchor = None
        for t in ("w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang"):
            anchor = rpr.find(qn(t))
            if anchor is not None:
                anchor.addprevious(h)
                return
        rpr.append(h)


def set_cell_text_style(cell, bold=False):
    for p in cell.paragraphs:
        for r in p.runs:
            if bold:
                r.font.bold = True


def column_weights(table):
    rows = table.rows
    ncols = len(rows[0].cells)
    weights = []
    for c in range(ncols):
        lens = []
        for r in rows:
            try:
                lens.append(len(r.cells[c].text.strip()))
            except IndexError:
                lens.append(0)
        avg = sum(lens) / max(1, len(lens))
        mx = max(lens)
        longest_word = 0
        for r in rows:
            try:
                for w in r.cells[c].text.split():
                    longest_word = max(longest_word, len(w))
            except IndexError:
                pass
        base = 0.55 * avg + 0.45 * min(mx, 70)
        weights.append(max(5.0, base, min(longest_word, 22) * 1.8 + 3))
    return weights


def style_table(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    # direction + centering + width
    for tag in ("w:bidiVisual", "w:jc", "w:tblW", "w:tblLayout", "w:tblBorders", "w:tblCellMar", "w:tblLook"):
        e = tblPr.find(qn(tag))
        if e is not None:
            tblPr.remove(e)
    total = int(MAX_W_IN * 1440)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(total))
    w.set(qn("w:type"), "dxa")
    tblPr.append(w)
    bv = OxmlElement("w:bidiVisual")
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    # schema order within tblPr: tblStyle, tblpPr, tblOverlap, bidiVisual, tblStyleRowBandSize, tblStyleColBandSize,
    # tblW, jc, tblCellSpacing, tblInd, tblBorders, shd, tblLayout, tblCellMar, tblLook
    style = tblPr.find(qn("w:tblStyle"))
    for el in (bv,):
        if style is not None:
            style.addnext(el)
        else:
            tblPr.insert(0, el)
    w_el = tblPr.find(qn("w:tblW"))
    w_el.addnext(jc)
    tblPr.append(lay)
    tblPr.append(mar)
    # grid + cell widths
    weights = column_weights(table)
    header = " ".join(c.text for c in table.rows[0].cells)
    if "קישור לרכישה" in header and len(weights) == 6:
        weights = [3, 17, 30, 5, 9, 30]  # BOM: leave room for real purchase URLs
    s = sum(weights)
    widths = [max(500, int(total * x / s)) for x in weights]
    widths[-1] += total - sum(widths)
    grid = tbl.tblGrid
    for gc, wd in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(wd))
    for ri, row in enumerate(table.rows):
        # header row repeat + no split rows
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
        if ri == 0:
            th = OxmlElement("w:tblHeader")
            trPr.append(th)
        for ci, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            for tag in ("w:tcW", "w:tcBorders", "w:shd", "w:vAlign"):
                e = tcPr.find(qn(tag))
                if e is not None:
                    tcPr.remove(e)
            tcw = OxmlElement("w:tcW")
            tcw.set(qn("w:w"), str(widths[ci] if ci < len(widths) else widths[-1]))
            tcw.set(qn("w:type"), "dxa")
            tcPr.insert(0, tcw)
            cell_borders(tcPr)
            if ri == 0:
                shade(tcPr, "E7ECE8")
            va = OxmlElement("w:vAlign")
            va.set(qn("w:val"), "center")
            tcPr.append(va)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else None
                if ri == 0:
                    for r in p.runs:
                        r.font.bold = True


def fit_size(path, max_w_in=MAX_W_IN, max_h_in=MAX_H_IN):
    with Image.open(path) as im:
        w, h = im.size
    # images are rendered at 2x CSS px; 96 css px per inch
    nat_w_in = (w / 2) / 96
    nat_h_in = (h / 2) / 96
    scale = min(1.0 if nat_w_in <= max_w_in else max_w_in / nat_w_in, max_h_in / nat_h_in if nat_h_in > max_h_in else 1.0)
    # upscale small images a bit (but never beyond max width) for legibility
    w_in = nat_w_in * scale
    h_in = nat_h_in * scale
    if w_in < max_w_in * 0.55 and h_in < max_h_in * 0.4:
        up = min(max_w_in * 0.75 / w_in, 1.6)
        w_in, h_in = w_in * up, h_in * up
    return Inches(w_in), Inches(h_in)


def replace_image_tokens(doc):
    for p in list(doc.paragraphs):
        t = p.text.strip()
        m = re.fullmatch(r"@@IMG\|(.*)\|(.*)@@", t)
        if not m:
            continue
        *files, caption = m.group(1).split("|") + [m.group(2)]
        # clear the token paragraph and put pictures in it (one picture per paragraph)
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        anchor = p
        for idx, f in enumerate(files):
            if idx == 0:
                target = p
            else:
                new_p = copy.deepcopy(p._p)
                for r in new_p.findall(qn("w:r")):
                    new_p.remove(r)
                anchor._p.addnext(new_p)
                from docx.text.paragraph import Paragraph
                target = Paragraph(new_p, p._parent)
            w, h = fit_size(f)
            target.add_run().add_picture(f, width=w, height=h)
            anchor = target
        cap = copy.deepcopy(p._p)
        for r in cap.findall(qn("w:r")):
            cap.remove(r)
        anchor._p.addnext(cap)
        from docx.text.paragraph import Paragraph
        cp = Paragraph(cap, p._parent)
        cp.style = doc.styles["Image Caption"]
        cp.paragraph_format.keep_with_next = False
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(12)
        cp.add_run(caption)
        # label bold (up to the colon)
        lab, _, rest = caption.partition(":")
        cp.runs[0].text = lab + ":"
        cp.runs[0].font.bold = True
        r2 = cp.add_run(rest)
        r2.font.bold = False


def style_special_paragraphs(doc):
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if re.match(r"^(איור|קוד) \d+:", txt) and p.style.name != "Image Caption":
            p.style = doc.styles["Image Caption"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = False
        elif re.match(r"^טבלה \d+:", txt):
            p.style = doc.styles["Image Caption"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.bold = True
        elif re.match(r"^\[\d+\]\s+[A-Za-z]", txt):
            # English bibliography entry: left-to-right paragraph, left aligned, hanging indent
            ppr = p._p.get_or_add_pPr()
            b = ppr.find(qn("w:bidi"))
            if b is None:
                b = OxmlElement("w:bidi")
                insert_ordered(ppr, b, PPR_AFTER_BIDI)
            b.set(qn("w:val"), "0")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Inches(0.4), Inches(-0.4)
            pf.line_spacing, pf.space_after = 1.15, Pt(6)
        elif txt.startswith("[הכנס"):
            ppr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            for side in ("top", "left", "bottom", "right"):
                e = OxmlElement(f"w:{side}")
                e.set(qn("w:val"), "dashed")
                e.set(qn("w:sz"), "8")
                e.set(qn("w:space"), "4")
                e.set(qn("w:color"), "B58900")
                bdr.append(e)
            insert_ordered(ppr, bdr, ["w:shd", "w:tabs", "w:suppressAutoHyphens", "w:bidi", "w:adjustRightInd",
                                      "w:snapToGrid", "w:spacing", "w:ind", "w:jc", "w:rPr"])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
    # headings: chapter headings start a new page
    for p in doc.paragraphs:
        if p.style.name == "Heading 1":
            p.paragraph_format.page_break_before = True


def post_runs(doc):
    def walk(paragraphs):
        for p in paragraphs:
            for r in p._p.iter(qn("w:r")):
                add_run_props(r)
            for run in p.runs:
                if PLACEHOLDER.search(run.text):
                    highlight_run(run)
                if HEBREW.search(run.text):
                    rpr = run._r.get_or_add_rPr()
                    if rpr.find(qn("w:rtl")) is None and run.style is not None and run.style.name != "Verbatim Char":
                        # w:rtl goes late in rPr; append is schema-safe enough before w:lang
                        lang = rpr.find(qn("w:lang"))
                        el = OxmlElement("w:rtl")
                        (lang.addprevious(el) if lang is not None else rpr.append(el))
    walk(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                walk(cell.paragraphs)


def add_page_number_footer(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, txt in (("begin", None), ("instr", " PAGE "), ("separate", None), ("text", "1"), ("end", None)):
        if kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            run._r.append(it)
        elif kind == "text":
            t = OxmlElement("w:t")
            t.text = txt
            run._r.append(t)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
    run.font.name = FONT
    run.font.size = Pt(11)


def para(doc_or_cell, text="", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6, hl=False,
         ltr=False):
    p = doc_or_cell.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    if not ltr:
        set_bidi(ppr)
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.15
    if text:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        rpr = r._r.get_or_add_rPr()
        set_fonts(rpr)
        if bold:
            rpr.append(OxmlElement("w:bCs"))
        cs = OxmlElement("w:szCs")
        cs.set(qn("w:val"), str(int(size * 2)))
        rpr.append(cs)
        if HEBREW.search(text):
            rpr.append(OxmlElement("w:rtl"))
        if hl:
            highlight_run(r)
    return p


def build_cover(doc):
    """Cover page in the style of the department's project template."""
    body = doc.element.body
    first = body[0]
    tmp = Document()  # scratch container to build paragraphs, then move them to the front
    cover = []

    def add(p):
        cover.append(p._p)

    add(para(tmp, "הפקולטה להנדסה", 16, True, after=0))
    add(para(tmp, "המחלקה להנדסת חשמל ומחשבים", 16, True, after=60))
    add(para(tmp, "פרויקט גמר", 22, True, after=6))
    add(para(tmp, "קורס חשיפה לחומרה", 14, True, after=50))
    add(para(tmp, "HomeGuard", 30, True, after=4, ltr=True))
    add(para(tmp, "מערכת אבטחה היברידית מרובת שכבות", 18, True, after=4))
    add(para(tmp, "מחשוב קצה על בקרי ESP32, בקנד Python עם זיהוי פנים ואפליקציית שליטה לטלפון נייד", 13, False,
             after=40))
    add(para(tmp, "תאריך הביצוע: dd/mm/yy", 14, True, after=4, hl=True))
    add(para(tmp, "תאריך ההגשה: dd/mm/yy", 14, True, after=40, hl=True))

    t = tmp.add_table(rows=5, cols=3)
    rows = [("", "שותף א'", "שותף ב'"),
            ("שמות המבצעים:", "גיא בר", "יאיר אברהם בונפיל"),
            ("מספר זהות:", "213263932", "325534121"),
            ('דוא"ל:', "[להשלים]@mail.tau.ac.il", "yairbonfil@mail.tau.ac.il"),
            ("חתימת המגישים:", "", "")]
    for ri, rv in enumerate(rows):
        for ci, val in enumerate(rv):
            cell = t.cell(ri, ci)
            cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
            pp = para(cell, val, 12, bold=(ci == 0 or ri == 0), align=WD_ALIGN_PARAGRAPH.CENTER, after=2,
                      hl=bool(PLACEHOLDER.search(val)), ltr=(ci > 0 and bool(re.search(r"[a-z]@", val))))
            if ri == 4:
                row = t.rows[ri]
                trpr = row._tr.get_or_add_trPr()
                h = OxmlElement("w:trHeight")
                h.set(qn("w:val"), "900")
                h.set(qn("w:hRule"), "atLeast")
                trpr.append(h)
    style_table(t)
    # style_table sizes to 6.3in and shades header; fine for the cover
    add_t = t._tbl
    cover.append(add_t)
    add(para(tmp, "", 6, after=14))
    add(para(tmp, "שמות המנחים: אריאלה צברי, עידן הבר", 13, True, after=2))
    add(para(tmp, "[לאמת מול הנחיית הפרויקט]", 11, False, after=0, hl=True))
    for el in cover:
        first.addprevious(el)


ORDERS = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
           "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection "
           "textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof "
           "snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText "
           "vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
             "tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "style": "name aliases basedOn next link autoRedefine hidden uiPriority semiHidden unhideWhenUsed qFormat locked "
             "personal personalCompose personalReply rsid pPr rPr tblPr trPr tcPr tblStylePr",
}
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def normalize_order(root):
    for local, order in ORDERS.items():
        rank = {f"{{{W_NS}}}{t}": i for i, t in enumerate(order.split())}
        for el in root.iter(f"{{{W_NS}}}{local}"):
            kids = list(el)
            for k in kids:
                if k.tail and k.tail.strip():
                    k.tail = None  # stray text between property elements is invalid
            if any(k.tag not in rank for k in kids):
                continue  # unknown child: leave alone
            srt = sorted(kids, key=lambda k: rank[k.tag])
            if srt != kids:
                for k in kids:
                    el.remove(k)
                for k in srt:
                    el.append(k)
    # math delimiters: begChr, sepChr, endChr, grow, shp, ctrlPr
    mrank = {f"{{{M_NS}}}{t}": i for i, t in enumerate("begChr sepChr endChr grow shp ctrlPr".split())}
    for el in root.iter(f"{{{M_NS}}}dPr"):
        kids = list(el)
        if all(k.tag in mrank for k in kids):
            srt = sorted(kids, key=lambda k: mrank[k.tag])
            if srt != kids:
                for k in kids:
                    el.remove(k)
                for k in srt:
                    el.append(k)


def fix_sections(doc):
    for sec in doc.element.body.iter(qn("w:sectPr")):
        pg = sec.find(qn("w:pgMar"))
        if pg is not None:
            for a, v in (("w:header", "708"), ("w:footer", "708"), ("w:gutter", "0")):
                if pg.get(qn(a)) is None:
                    pg.set(qn(a), v)


def finalize_doc(raw: Path, out: Path):
    doc = Document(str(raw))
    replace_image_tokens(doc)
    style_special_paragraphs(doc)
    for t in doc.tables:
        style_table(t)
    post_runs(doc)
    build_cover(doc)
    add_page_number_footer(doc)
    cp = doc.core_properties
    cp.title = "HomeGuard — מערכת אבטחה היברידית מרובת שכבות"
    cp.author = ""
    cp.language = "he-IL"
    # document-level default language and bidi
    st = doc.styles.element
    dd = st.find(qn("w:docDefaults"))
    if dd is not None:
        rpr = dd.find(qn("w:rPrDefault")).find(qn("w:rPr"))
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rpr.append(lang)
        lang.set(qn("w:val"), "en-US")
        lang.set(qn("w:bidi"), "he-IL")
    fix_sections(doc)
    normalize_order(doc.element)
    normalize_order(doc.styles.element)
    for sec in doc.sections:
        for part in (sec.footer, sec.first_page_footer):
            normalize_order(part._element)
    doc.save(str(out))


def patch_package(path: Path):
    """Pandoc's stock parts have schema nits Word tolerates but the validator flags (short w:nsid values, an
    out-of-order settings.xml). Fix them so the file is clean."""
    import zipfile
    tmp_path = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/numbering.xml":
                txt = data.decode("utf-8")
                txt = re.sub(r'(<w:nsid w:val=")([0-9A-Fa-f]+)(")', lambda m: m.group(1) + m.group(2).zfill(8) + m.group(3), txt)
                data = txt.encode("utf-8")
            elif item.filename == "word/settings.xml":
                data = (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                    '<w:settings xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:defaultTabStop w:val="720"/><w:characterSpacingControl w:val="doNotCompress"/>'
                    '<w:compat><w:compatSetting w:name="compatibilityMode" '
                    'w:uri="http://schemas.microsoft.com/office/word" w:val="15"/></w:compat>'
                    '</w:settings>'
                ).encode("utf-8")
            zout.writestr(item, data)
    tmp_path.replace(path)


def main():
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        ref = tmp / "ref.docx"
        make_reference(ref)
        master = build_master(tmp)
        raw = tmp / "raw.docx"
        subprocess.run(["pandoc", str(master), "-f", "markdown-smart", "-t", "docx",
                        "--shift-heading-level-by=-1", f"--reference-doc={ref}", "-o", str(raw)], check=True)
        finalize_doc(raw, OUT)
        patch_package(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
